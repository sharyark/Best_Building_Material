"""Render a ReportTable as Excel (openpyxl) or PDF (WeasyPrint)."""

import html
import io
from decimal import Decimal

from fastapi import Response
from fastapi.responses import JSONResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from weasyprint import HTML

from app.core.config import get_settings
from app.modules.reports.table import ReportTable, fmt_cell

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NUMERIC = {"money": "#,##0.00", "qty": "#,##0.###", "int": "#,##0", "drcr": "#,##0.00;[Red]-#,##0.00", "percent": "0.0"}


def _filename(table: ReportTable, ext: str) -> str:
    safe = "".join(c if c.isalnum() else "_" for c in table.title).strip("_").lower()
    return f"{safe}.{ext}"


def to_xlsx(table: ReportTable) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = table.title[:31]
    ws.append([table.title])
    ws["A1"].font = Font(bold=True, size=14)
    if table.subtitle:
        ws.append([table.subtitle])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([c.label for c in table.columns])
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")

    def write(row: dict, bold: bool = False) -> None:
        values = []
        for c in table.columns:
            v = row.get(c.key)
            if c.type in NUMERIC and v not in (None, ""):
                v = float(Decimal(str(v)))
            elif c.type == "date" and v:
                v = fmt_cell(v, "date")
            values.append(v)
        ws.append(values)
        for cell, c in zip(ws[ws.max_row], table.columns, strict=True):
            if c.type in NUMERIC:
                cell.number_format = NUMERIC[c.type]
            if bold:
                cell.font = Font(bold=True)

    for row in table.rows:
        write(row)
    if table.totals:
        write({"_": None, **table.totals}, bold=True)
    for idx, col in enumerate(table.columns, start=1):
        widest = max([len(col.label)] + [len(fmt_cell(r.get(col.key), col.type)) for r in table.rows[:200]])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(widest + 2, 10), 50)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


PDF_CSS = """
@page { size: A4 %s; margin: 14mm 12mm; @bottom-right { content: "Page " counter(page) " of " counter(pages); font-size: 8pt; color: #666; } }
body { font-family: DejaVu Sans, Arial, sans-serif; font-size: 9pt; color: #111; }
h1 { font-size: 15pt; margin: 0 0 2px; } h2 { font-size: 10pt; margin: 0 0 8px; color: #555; font-weight: normal; }
.biz { font-size: 9pt; color: #555; margin-bottom: 6px; }
table { width: 100%%; border-collapse: collapse; }
th { background: #1f4e79; color: #fff; padding: 4px 6px; text-align: left; font-size: 8.5pt; }
td { padding: 3px 6px; border-bottom: 1px solid #ddd; vertical-align: top; }
td.num, th.num { text-align: right; white-space: nowrap; }
tr.total td { font-weight: bold; border-top: 2px solid #333; background: #f3f3f3; }
.meta { margin-top: 10px; font-size: 8.5pt; color: #444; }
"""

NUMERIC_TYPES = {"money", "qty", "int", "drcr", "percent"}


def to_pdf(table: ReportTable) -> bytes:
    esc = html.escape
    head = "".join(f'<th class="{"num" if c.type in NUMERIC_TYPES else ""}">{esc(c.label)}</th>' for c in table.columns)

    def tr(row: dict, cls: str = "") -> str:
        tds = "".join(
            f'<td class="{"num" if c.type in NUMERIC_TYPES else ""}">{esc(fmt_cell(row.get(c.key), c.type))}</td>'
            for c in table.columns
        )
        return f'<tr class="{cls}">{tds}</tr>'

    body = "".join(tr(r) for r in table.rows)
    if table.totals:
        body += tr(table.totals, "total")
    orientation = "landscape" if len(table.columns) > 6 else "portrait"
    subtitle = f"<h2>{esc(table.subtitle)}</h2>" if table.subtitle else ""
    doc = (
        f"<html><head><meta charset='utf-8'><style>{PDF_CSS % orientation}</style></head><body>"
        f"<div class='biz'>{esc(get_settings().business_name)}</div><h1>{esc(table.title)}</h1>{subtitle}"
        f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></body></html>"
    )
    return HTML(string=doc).write_pdf()


def respond(table: ReportTable, fmt: str) -> Response:
    if fmt == "xlsx":
        return Response(
            to_xlsx(table), media_type=XLSX_TYPE,
            headers={"Content-Disposition": f'attachment; filename="{_filename(table, "xlsx")}"'},
        )  # fmt: skip
    if fmt == "pdf":
        return Response(
            to_pdf(table), media_type="application/pdf",
            headers={"Content-Disposition": f'inline; filename="{_filename(table, "pdf")}"'},
        )  # fmt: skip
    return JSONResponse(table.to_dict())
